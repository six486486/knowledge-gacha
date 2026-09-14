"""Synthetic, local-only document-import fixtures. No user documents or model calls."""
from pathlib import Path
from io import BytesIO
import argparse
import random
import zipfile
import struct
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader
from pypdf import PdfReader, PdfWriter
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

root = Path(__file__).resolve().parents[1]
fixtures = root / "tests" / "fixtures" / "documents"
fixtures.mkdir(parents=True, exist_ok=True)
font_path = Path("C:/Windows/Fonts/simhei.ttf")
pdfmetrics.registerFont(TTFont("FixtureChinese", str(font_path)))
COST = "每多生产一个单位额外增加的成本，称为边际成本。比较新增收入与新增支出，才能判断是否值得增加产量。"
CACHE = "缓存有效期决定一份数据可以复用多久。有效期较短时更容易保持新鲜，但会增加回源请求。"
EXCLUDED = "EXCLUDED_RANGE_TOKEN：这部分是未选内容，只能留在本地，不能进入制卡请求。"

def line(c, text, x, y, size=12):
    c.setFont("FixtureChinese", size)
    c.drawString(x, y, text)

def body(c, text, x, y, width=38, size=12, leading=20):
    for start in range(0, len(text), width):
        line(c, text[start:start + width], x, y, size)
        y -= leading

def page(c, number, title, text):
    line(c, "导入回归样本 · 固定页眉", 45, 819, 9)
    line(c, str(number), 286, 20, 9)
    line(c, title, 45, 750, 18)
    body(c, text, 45, 704)

c = canvas.Canvas(str(fixtures / "chinese.pdf"), pagesize=(595, 842), invariant=1)
for number, (title, text) in enumerate([("边际成本", COST * 3), ("缓存策略", CACHE * 3), ("未选章节", EXCLUDED)], 1):
    c.bookmarkPage(f"chapter{number}")
    c.addOutlineEntry(title, f"chapter{number}", 0)
    page(c, number, title, text)
    c.showPage()
c.save()

c = canvas.Canvas(str(fixtures / "two-column.pdf"), pagesize=(595, 842), invariant=1)
line(c, "双栏阅读顺序测试", 45, 795, 18)
line(c, "左栏：缓存", 45, 750, 13)
body(c, "左栏开始。" + CACHE * 4 + "左栏结束。", 45, 720, 20, 10, 17)
line(c, "右栏：成本", 320, 750, 13)
body(c, "右栏开始。" + COST * 4 + "右栏结束。", 320, 720, 20, 10, 17)
c.save()

scan = Image.new("RGB", (1190, 1684), "white")
draw = ImageDraw.Draw(scan)
draw.text((90, 130), "扫描页：边际成本", font=ImageFont.truetype(str(font_path), 42), fill="black")
for index, text in enumerate([COST[:24], COST[24:], "识别后可以修改文字，再选择范围制卡。"]):
    draw.text((90, 260 + index * 70), text, font=ImageFont.truetype(str(font_path), 30), fill="black")
image = ImageReader(scan)
for name, mixed in [("scanned.pdf", False), ("mixed.pdf", True)]:
    c = canvas.Canvas(str(fixtures / name), pagesize=(595, 842), invariant=1)
    if mixed:
        page(c, 1, "可直接提取的正文", COST)
        c.showPage()
    c.drawImage(image, 0, 0, width=595, height=842)
    c.showPage()
    c.save()

writer = PdfWriter(clone_from=str(fixtures / "chinese.pdf"))
writer.encrypt("fixture-password")
with (fixtures / "encrypted.pdf").open("wb") as stream:
    writer.write(stream)
(fixtures / "corrupt.pdf").write_bytes(b"%PDF-1.7\nThis is a deliberately corrupt test file.\n%%EOF")

doc = Document()
doc.add_heading("边际成本", 1)
doc.add_paragraph(COST)
doc.add_paragraph("先比较新增收入。", "List Number")
doc.add_paragraph("再扣除新增支出。", "List Number")
doc.add_heading("缓存策略", 1)
doc.add_paragraph(CACHE)
table = doc.add_table(rows=3, cols=2)
for row, values in zip(table.rows, [("设置", "含义"), ("短有效期", "更快更新缓存，但回源更多。"), ("长有效期", "减少回源，但旧数据可能保留较久。")]):
    for cell, value in zip(row.cells, values):
        cell.text = value
doc.add_heading("不选择的章节", 1)
doc.add_paragraph(EXCLUDED)
doc.add_paragraph('<script>window.DOCX_INJECTION = true</script> 必须作为文字显示。')
doc.save(fixtures / "headings-lists-table.docx")

(fixtures / "corrupt.docx").write_bytes(b"PK\x03\x04corrupt fixture")
with zipfile.ZipFile(fixtures / "oversized-entry.docx", "w", compression=zipfile.ZIP_DEFLATED) as archive:
    archive.writestr("word/document.xml", b"x" * (16 * 1024 * 1024 + 1))
with zipfile.ZipFile(fixtures / "forged-entry.docx", "w", compression=zipfile.ZIP_DEFLATED) as archive:
    archive.writestr("word/document.xml", b"x" * (1024 * 1024))
forged = bytearray((fixtures / "forged-entry.docx").read_bytes())
central = forged.index(b"PK\x01\x02")
struct.pack_into("<I", forged, central + 24, 1)
struct.pack_into("<I", forged, 22, 1)
(fixtures / "forged-entry.docx").write_bytes(forged)

parser = argparse.ArgumentParser()
parser.add_argument("--benchmark", action="store_true")
args = parser.parse_args()
if args.benchmark:
    output = root.parent / "test-results" / "import-bench"
    output.mkdir(parents=True, exist_ok=True)
    target = output / "limit-20mb-101pages.pdf"
    raw = BytesIO()
    c = canvas.Canvas(raw, pagesize=(595, 842), invariant=1)
    rng = random.Random(4)
    for number in range(1, 102):
        page(c, number, f"第 {number} 页负载样本", COST * 6)
        noise = Image.frombytes("RGB", (350, 350), rng.randbytes(350 * 350 * 3))
        jpeg = BytesIO()
        noise.save(jpeg, format="JPEG", quality=90)
        c.drawImage(ImageReader(jpeg), 100, 60, width=300, height=300)
        c.showPage()
    c.save()
    writer = PdfWriter(clone_from=BytesIO(raw.getvalue()))
    budget = 20 * 1024 * 1024
    padding = max(0, budget - len(raw.getvalue()) - 100000)
    writer.add_attachment("synthetic-padding.bin", rng.randbytes(padding))
    with target.open("wb") as stream:
        writer.write(stream)
    # The benchmark records the actual size; never claim a larger limit was tested.
    print(f"Benchmark: {target.name}, {target.stat().st_size} bytes, 101 pages")
print("Created Chinese, column, scan, mixed, encrypted/corrupt PDF and structured DOCX fixtures")
